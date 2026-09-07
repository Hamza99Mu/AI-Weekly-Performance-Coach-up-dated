import json
import os
import re
from io import BytesIO
from typing import Any, Dict, Optional

from dotenv import load_dotenv
from docx import Document
from groq import Groq
from prompts import SYSTEM_PROMPT, build_analysis_prompt

load_dotenv()

REQUIRED = {
    "weekly_analysis", "time_saved", "time_eaters",
    "weekly_audit", "next_week_plan", "final_result"
}

def extract_docx_text(file_or_bytes) -> str:
    if hasattr(file_or_bytes, "read"):
        data = file_or_bytes.read()
    elif isinstance(file_or_bytes, (bytes, bytearray)):
        data = bytes(file_or_bytes)
    else:
        raise TypeError("DOCX input must be an uploaded file or bytes.")

    doc = Document(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)

def _api_key(api_key: Optional[str]) -> str:
    key = api_key or os.getenv("GROQ_API_KEY")
    if not key:
        raise ValueError("Groq API key not found. Add GROQ_API_KEY to .env or enter it in the sidebar.")
    return key

def _parse_json(text: str) -> Dict[str, Any]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s*```$", "", cleaned).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start >= 0 and end > start:
            try:
                return json.loads(cleaned[start:end + 1])
            except json.JSONDecodeError as e:
                raise ValueError(f"Groq returned malformed JSON: {e}") from e
        raise ValueError("Groq did not return valid JSON.")

def _validate(result: Dict[str, Any]) -> None:
    missing = REQUIRED - set(result)
    if missing:
        raise ValueError("AI response is missing: " + ", ".join(sorted(missing)))
    analysis = result["weekly_analysis"]
    score = analysis.get("overall_score")
    if not isinstance(score, (int, float)) or not 0 <= score <= 100:
        raise ValueError("Overall score must be between 0 and 100.")
    for field in ("total_tracked_hours", "high_impact_hours", "low_value_hours", "goal_completion_percent"):
        value = analysis.get(field)
        if not isinstance(value, (int, float)) or value < 0:
            raise ValueError(f"{field} must be a non-negative number.")
    if analysis["goal_completion_percent"] > 100:
        raise ValueError("Goal completion cannot exceed 100%.")

def analyze_week(weekly_text: str, goals_text: str = "", api_key: Optional[str] = None) -> Dict[str, Any]:
    if not weekly_text.strip():
        raise ValueError("Weekly activities cannot be empty.")

    client = Groq(api_key=_api_key(api_key))
    model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")

    response = client.chat.completions.create(
        model=model,
        temperature=0.2,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": build_analysis_prompt(weekly_text.strip(), goals_text.strip())},
        ],
    )

    content = response.choices[0].message.content or ""
    result = _parse_json(content)
    _validate(result)
    return result
